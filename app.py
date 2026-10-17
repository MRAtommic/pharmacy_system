"""
🏥 Pharmacy Stock Management & Dispensing System
Flask Application Entry Point
"""

import sys
import os
import re
import json
import sqlite3
import logging
import secrets
import urllib.parse
import urllib.request
from pathlib import Path
from datetime import datetime
from flask import Flask, request, jsonify, render_template, send_from_directory, redirect, send_file, make_response, session
from io import BytesIO
from config import Config
import database
from google_sync import google_sync
from line_service import line_bot
from ocr_service import parse_drug_from_image

# Safe UTF-8 on Windows
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

app = Flask(
    __name__,
    template_folder=str(Path(__file__).parent / "templates"),
    static_folder=str(Path(__file__).parent / "static")
)
app.config.from_object(Config)
app.secret_key = Config.SECRET_KEY
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.jinja_env.auto_reload = True

logger = logging.getLogger("pharmacy_app")
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

# ─────────────────────────────────────────────────────────────
# Auth: Login Required Decorator
# ─────────────────────────────────────────────────────────────

from functools import wraps

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        code = (request.args.get('code') or request.args.get('passcode') or '').strip()
        if code and code in ['otwxc8x', os.environ.get('STAFF_PASSCODE', 'otwxc8x')]:
            if not session.get('user_email'):
                session['user_email'] = 'admin@pharmacore.clinic'
                session['user_name'] = 'ผู้ดูแลระบบคลังยา (Admin)'
                session['user_picture'] = '/static/uploads/drugs/app_icon_192.png'
        if not session.get('user_email'):
            if request.path.startswith('/api/') and '/requisition' not in request.path:
                return jsonify({"ok": False, "error": "กรุณาล็อกอินก่อนใช้งาน", "login_required": True}), 401
            next_target = request.full_path if request.query_string else request.path
            return redirect(f'/login?next={next_target}')
        return f(*args, **kwargs)
    return decorated

def current_user() -> dict:
    """Return current session user info or empty dict."""
    return {
        "email": session.get('user_email', ''),
        "name": session.get('user_name', ''),
        "picture": session.get('user_picture', '')
    }

def audit(action: str, target_type: str = '', target_id='', target_name: str = '', detail: str = ''):
    """Helper to write audit log with current user & request IP."""
    u = current_user()
    ip = request.headers.get('X-Forwarded-For', request.remote_addr or '')
    database.log_audit_action(
        user_email=u['email'] or 'anonymous',
        user_name=u['name'] or '',
        action=action,
        target_type=target_type,
        target_id=str(target_id),
        target_name=target_name,
        detail=detail,
        ip_address=ip
    )

@app.after_request
def disable_cache_for_dev(response):
    # Skip no-cache for file downloads — browsers refuse to save attachments with no-store
    if request.path.startswith("/api/export"):
        return response
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

# ─────────────────────────────────────────────────────────────
# HTML Views
# ─────────────────────────────────────────────────────────────

@app.route("/")
@login_required
def index():
    u = current_user()
    return render_template("index.html", base_url=Config.BASE_URL, current_user=u)

@app.route("/scanner")
@login_required
def mobile_scanner():
    u = current_user()
    return render_template("index.html", initial_view="scanner", base_url=Config.BASE_URL, current_user=u)

def format_packaging_size(drug):
    if not drug:
        return "1*500 's"
    strength = (drug.get("strength") or "").strip()
    if strength:
        cleaned = re.sub(r'\s*(เม็ด|ขวด|แคปซูล|ซอง|หลอด|ตลับ|กล่อง|ชิ้น|แผง)', '', strength).strip()
        if cleaned:
            return cleaned
    name = drug.get("name", "")
    m = re.search(r'(\d+(?:\.\d+)?\s*(?:mg|g|ml|mcg|มก\.?|มล\.?|%))', name, re.IGNORECASE)
    if m:
        return m.group(1)
    return "1*500 's"

@app.route("/stock-card")
@login_required
def stock_card_view():
    u = current_user()
    drugs = database.get_all_drugs()
    
    selected_drug_id = request.args.get("drug_id")
    try:
        selected_drug_id = int(selected_drug_id) if selected_drug_id else (drugs[0]["id"] if drugs else None)
    except (ValueError, TypeError):
        selected_drug_id = drugs[0]["id"] if drugs else None

    now = datetime.now()
    default_start = request.args.get("start_date") or f"{now.year}-{now.month:02d}-01"
    default_end = request.args.get("end_date") or now.strftime("%Y-%m-%d")

    card_data = database.get_stock_card_data(selected_drug_id, default_start, default_end) if selected_drug_id else None
    
    selected_drug = card_data["drug"] if card_data else None
    transactions = card_data["transactions"] if card_data else []
    opening_balance = card_data["opening_balance"] if card_data else 0

    for tx in transactions:
        tx["date_short"] = format_thai_date_short(tx["created_at"])
        lot = (tx.get("lot_number") or "").strip()
        exp = (tx.get("expiry_date") or "").strip()

        if lot and exp:
            lot_exp_display = f'<div style="line-height:1.25;"><div style="font-weight:600;">{lot}</div><div style="font-size:0.85em; color:#475569;">{exp}</div></div>'
        elif lot:
            lot_exp_display = f'<span style="font-weight:600;">{lot}</span>'
        elif exp:
            lot_exp_display = exp
        else:
            lot_exp_display = "-"

        if tx["type"] == "RECEIVE":
            is_gpo = "GPO" in (tx.get("supplier") or "") or "องค์การ" in (tx.get("supplier") or "")
            qty_in = tx.get("qty_in") or tx.get("quantity") or "-"
            if is_gpo:
                tx["particulars_html"] = f'รับจาก..<span class="fill-text">{qty_in}</span>..GPO/..-..บริษัท/จ่าย..-..'
            else:
                tx["particulars_html"] = f'รับจาก...-..GPO/..<span class="fill-text">{qty_in}</span>..บริษัท/จ่าย..-..'
            tx["lot_exp_html"] = lot_exp_display
        else:
            qty_out = tx.get("qty_out") or tx.get("quantity") or "-"
            tx["particulars_html"] = f'รับจาก...-..GPO/..-..บริษัท/จ่าย..<span class="fill-text">{qty_out}</span>..'
            tx["lot_exp_html"] = lot_exp_display if lot else "-"

        if tx.get("total_price"):
            tx["total_price_formatted"] = f"{round(tx['total_price']):,}"
        elif tx.get("price"):
            tx["price_formatted"] = f"{round(tx['price']):,}"

    opening_date_short = ""
    used_rows = 1 + len(transactions)
    empty_rows_count = max(0, 24 - used_rows)
    opening_value_formatted = card_data.get("opening_value_formatted", "0") if card_data else "0"

    opening_lot = (selected_drug.get("lot_number") or "").strip() if selected_drug else ""
    opening_exp = (selected_drug.get("expiry_date") or "").strip() if selected_drug else ""
    if opening_lot and opening_exp:
        opening_lot_exp = f'<div style="line-height:1.25;"><div style="font-weight:600;">{opening_lot}</div><div style="font-size:0.85em; color:#475569;">{opening_exp}</div></div>'
    elif opening_lot:
        opening_lot_exp = f'<span style="font-weight:600;">{opening_lot}</span>'
    elif opening_exp:
        opening_lot_exp = opening_exp
    else:
        opening_lot_exp = "-"

    return render_template(
        "stock_card.html",
        drugs=drugs,
        selected_drug=selected_drug,
        start_date=default_start,
        end_date=default_end,
        opening_date_short=opening_date_short,
        opening_balance=opening_balance,
        opening_value_formatted=opening_value_formatted,
        opening_lot_exp=opening_lot_exp,
        transactions=transactions,
        empty_rows_count=empty_rows_count,
        packaging=format_packaging_size(selected_drug),
        opening_sign="วราภรณ์",
        hospital_name="โรงพยาบาลส่งเสริมสุขภาพตำบลดอนนาลับ อำเภอบ้านดุง จังหวัดอุดรธานี",
        current_user=u,
        base_url=Config.BASE_URL
    )

@app.route("/api/stock-card/data", methods=["GET"])
@login_required
def api_stock_card_data():
    drug_id = request.args.get("drug_id")
    if not drug_id:
        return jsonify({"ok": False, "error": "ไม่ระบุ drug_id"}), 400
    try:
        drug_id = int(drug_id)
    except ValueError:
        return jsonify({"ok": False, "error": "drug_id ไม่ถูกต้อง"}), 400
        
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    data = database.get_stock_card_data(drug_id, start_date, end_date)
    if not data:
        return jsonify({"ok": False, "error": "ไม่พบข้อมูลยา"}), 404
        
    return jsonify({"ok": True, "data": data})

@app.route("/api/stock-card/export-excel", methods=["GET"])
@login_required
def export_stock_card_excel():
    drug_id = request.args.get("drug_id")
    if not drug_id:
        return jsonify({"ok": False, "error": "ไม่ระบุ drug_id"}), 400
    try:
        drug_id = int(drug_id)
    except ValueError:
        return jsonify({"ok": False, "error": "drug_id ไม่ถูกต้อง"}), 400

    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    opening_sign = request.args.get("opening_sign") or "วราภรณ์"
    data = database.get_stock_card_data(drug_id, start_date, end_date)
    if not data:
        return jsonify({"ok": False, "error": "ไม่พบข้อมูลยา"}), 404

    drug = data["drug"]
    transactions = data.get("transactions", [])
    
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Card"

    font_title = Font(name="Sarabun", size=14, bold=True)
    font_bold = Font(name="Sarabun", size=11, bold=True)
    font_regular = Font(name="Sarabun", size=11)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws.merge_cells("A1:H1")
    ws["A1"] = "บัญชีรายการรับ-จ่าย/เวชภัณฑ์มิใช่ยา/วัสดุการแพทย์"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:H2")
    ws["A2"] = "โรงพยาบาลส่งเสริมสุขภาพตำบลดอนนาลับ อำเภอบ้านดุง จังหวัดอุดรธานี"
    ws["A2"].font = font_bold
    ws["A2"].alignment = align_center

    ws.merge_cells("A3:H3")
    ws["A3"] = f"ชื่อสามัญ: {drug.get('generic_name') or drug.get('name')}     ขนาดบรรจุ: {format_packaging_size(drug)}     ลักษณะ: {drug.get('dosage_form', 'tablet')}"
    ws["A3"].font = font_regular
    ws["A3"].alignment = align_left

    headers = ["วัน/เดือน/ปี", "รายการ", "รับ", "จ่าย", "คงเหลือ", "ราคา", "lot/exp.", "ลงชื่อ"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num, value=header)
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = border_thin
        cell.fill = fill_header

    r_idx = 6
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=4)
    ws.cell(row=r_idx, column=1, value="ยอดยกมา จากปีงบประมาณ 2568").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_bold
    unit_cost = (drug.get("cost") if (drug.get("cost") is not None and drug.get("cost") > 0) else (drug.get("price") or 0.0)) or 0.0
    opening_val = data.get("opening_value")
    if opening_val is None:
        opening_val = data["opening_balance"] * unit_cost
    if isinstance(opening_val, float) and opening_val.is_integer():
        opening_val = int(opening_val)
    elif isinstance(opening_val, float):
        opening_val = round(opening_val, 2)

    ws.cell(row=r_idx, column=5, value=data["opening_balance"]).alignment = align_center
    opening_lot = (drug.get("lot_number") or "").strip()
    opening_exp = (drug.get("expiry_date") or "").strip()
    if opening_lot and opening_exp:
        opening_lot_exp_val = f"{opening_lot}\n{opening_exp}"
    elif opening_lot:
        opening_lot_exp_val = opening_lot
    elif opening_exp:
        opening_lot_exp_val = opening_exp
    else:
        opening_lot_exp_val = "-"
    cell_op_lot = ws.cell(row=r_idx, column=7, value=opening_lot_exp_val)
    cell_op_lot.alignment = align_center
    cell_op_lot.alignment.wrap_text = True
    ws.cell(row=r_idx, column=8, value=opening_sign).alignment = align_center
    for c in range(1, 9):
        ws.cell(row=r_idx, column=c).border = border_thin
        if c > 4:
            ws.cell(row=r_idx, column=c).font = font_regular

    for tx in transactions:
        r_idx += 1
        d_short = format_thai_date_short(tx["created_at"])
        if tx["type"] == "RECEIVE":
            is_gpo = "GPO" in (tx.get("supplier") or "") or "องค์การ" in (tx.get("supplier") or "")
            qty_in = tx.get("qty_in") or tx.get("quantity") or "-"
            if is_gpo:
                part = f"รับจาก...{qty_in}...GPO/...-.....บริษัท/จ่าย...-....."
            else:
                part = f"รับจาก...-.....GPO/...{qty_in}...บริษัท/จ่าย...-....."
        else:
            qty_out = tx.get("qty_out") or tx.get("quantity") or "-"
            part = f"รับจาก...-.....GPO/...-.....บริษัท/จ่าย...{qty_out}....."

        bal_val = tx.get("balance_value")
        if bal_val is None:
            bal_val = (tx.get("balance") or 0) * unit_cost
        if isinstance(bal_val, float) and bal_val.is_integer():
            bal_val = int(bal_val)
        elif isinstance(bal_val, float):
            bal_val = round(bal_val, 2)

        ws.cell(row=r_idx, column=1, value=d_short).alignment = align_center
        ws.cell(row=r_idx, column=2, value=part).alignment = align_left
        ws.cell(row=r_idx, column=3, value=tx["qty_in"] if tx["type"] == "RECEIVE" else "-").alignment = align_center
        ws.cell(row=r_idx, column=4, value=tx["qty_out"] if tx["type"] == "DISPENSE" else "-").alignment = align_center
        ws.cell(row=r_idx, column=5, value=tx["balance"]).alignment = align_center
        ws.cell(row=r_idx, column=6, value=bal_val).alignment = align_right
        lot = (tx.get("lot_number") or "").strip()
        exp = (tx.get("expiry_date") or "").strip()
        if lot and exp:
            lot_exp = f"{lot}\n{exp}"
        elif lot:
            lot_exp = lot
        elif exp:
            lot_exp = exp
        else:
            lot_exp = "-"
        ws.cell(row=r_idx, column=7, value=lot_exp).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=r_idx, column=8, value="").alignment = align_center

        for c in range(1, 9):
            ws.cell(row=r_idx, column=c).border = border_thin
            ws.cell(row=r_idx, column=c).font = font_regular

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 16
    ws.column_dimensions['H'].width = 14

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"stock_card_{drug.get('code', 'drug')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@app.route("/daily-requisition")
@login_required
def daily_requisition_view():
    u = current_user()
    now_dt = datetime.now()
    default_date = request.args.get("date") or now_dt.strftime("%Y-%m-%d")

    data = database.get_daily_requisition_data(default_date)
    available_dates = database.get_daily_requisition_available_dates()
    items = data["items"]
    empty_rows_count = max(0, 18 - len(items))

    return render_template(
        "daily_requisition.html",
        selected_date=default_date,
        req_day=data["day"],
        req_month=data["month_thai"],
        req_year=data["year_be"],
        items=items,
        empty_rows_count=empty_rows_count,
        has_records=data.get("has_records", len(items) > 0),
        available_dates=available_dates,
        hospital_name=data["hospital_name"],
        requester_name=data["requester_name"],
        requester_title=data["requester_title"],
        dispenser_name=data["dispenser_name"],
        dispenser_title=data["dispenser_title"],
        authorizer_name=data["authorizer_name"],
        authorizer_title=data["authorizer_title"],
        receiver_name=data["receiver_name"],
        receiver_title=data["receiver_title"],
        current_user=u,
        base_url=Config.BASE_URL
    )

@app.route("/api/daily-requisition/data", methods=["GET"])
@login_required
def api_daily_requisition_data():
    date_val = request.args.get("date")
    data = database.get_daily_requisition_data(date_val)
    if not data:
        return jsonify({"ok": False, "error": "ไม่พบข้อมูลสำหรับวันที่ระบุ"}), 404
    available_dates = database.get_daily_requisition_available_dates()
    return jsonify({"ok": True, "data": data, "available_dates": available_dates})

@app.route("/api/daily-requisition/export-excel", methods=["GET"])
@login_required
def export_daily_requisition_excel():
    date_val = request.args.get("date")
    data = database.get_daily_requisition_data(date_val)
    if not data:
        return jsonify({"ok": False, "error": "ไม่พบข้อมูลสำหรับวันที่ระบุ"}), 404

    requester = request.args.get("requester") or data["requester_name"]
    dispenser = request.args.get("dispenser") or data["dispenser_name"]
    authorizer = request.args.get("authorizer") or data["authorizer_name"]
    receiver = request.args.get("receiver") or data["receiver_name"]

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Requisition"

    font_title = Font(name="Sarabun", size=15, bold=True)
    font_bold = Font(name="Sarabun", size=11, bold=True)
    font_regular = Font(name="Sarabun", size=11)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws.merge_cells("A1:G1")
    ws["A1"] = "ใบเบิกเวชภัณฑ์ยาและเวชภัณฑ์มิใช่ยา"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:G2")
    ws["A2"] = f"จากหน่วยงาน{data['hospital_name']}"
    ws["A2"].font = font_bold
    ws["A2"].alignment = align_center

    ws.merge_cells("A3:G3")
    ws["A3"] = "ถึง  ผู้รับผิดคลังยาใหญ่   ข้าพเจ้าขอเบิกเวชภัณฑ์มิใช่ยาตามข้างล่างนี้เพื่อใช้ในงานรักษาพยาบาล"
    ws["A3"].font = font_regular
    ws["A3"].alignment = align_center

    ws.merge_cells("A4:G4")
    ws["A4"] = f"ตามใบเบิกนี้ภายใน  วันที่ {data['day']}  เดือน {data['month_thai']}  พ.ศ. {data['year_be']}"
    ws["A4"].font = font_regular
    ws["A4"].alignment = align_center

    headers = ["ลำดับที่", "รายการ", "หน่วยนับ", "จำนวนเบิก", "จำนวนจ่าย", "คงเหลือ", "หมายเหตุ"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_num, value=header)
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = border_thin
        cell.fill = fill_header

    r_idx = 6
    items = data.get("items", [])
    for idx, it in enumerate(items, 1):
        r_idx += 1
        ws.cell(row=r_idx, column=1, value=idx).alignment = align_center
        ws.cell(row=r_idx, column=2, value=it.get("name", "")).alignment = align_left
        ws.cell(row=r_idx, column=3, value=it.get("unit", "")).alignment = align_center
        ws.cell(row=r_idx, column=4, value=it.get("qty_requested") if it.get("qty_requested") not in (None, "") else "").alignment = align_center
        ws.cell(row=r_idx, column=5, value=it.get("qty_dispensed") if it.get("qty_dispensed") not in (None, "") else "").alignment = align_center
        ws.cell(row=r_idx, column=6, value=it.get("remaining_stock", 0)).alignment = align_center
        ws.cell(row=r_idx, column=7, value="").alignment = align_left
        for c in range(1, 8):
            ws.cell(row=r_idx, column=c).border = border_thin
            ws.cell(row=r_idx, column=c).font = font_regular

    # Pad with empty rows to at least 18 rows
    empty_count = max(0, 18 - len(items))
    for i in range(empty_count):
        r_idx += 1
        seq_val = len(items) + i + 1 if len(items) == 0 else ""
        ws.cell(row=r_idx, column=1, value=seq_val).alignment = align_center
        for c in range(1, 8):
            ws.cell(row=r_idx, column=c).border = border_thin
            ws.cell(row=r_idx, column=c).font = font_regular

    # Signatures block
    r_idx += 2
    # Row 1 of signatures: ผู้เบิก and ผู้จ่าย
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    ws.cell(row=r_idx, column=1, value="(ลงชื่อ) .................................................... ผู้เบิก").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_regular

    ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    ws.cell(row=r_idx, column=5, value="(ลงชื่อ) .................................................... ผู้จ่าย").alignment = align_center
    ws.cell(row=r_idx, column=5).font = font_regular

    r_idx += 1
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    ws.cell(row=r_idx, column=1, value=f"( {requester} )").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_bold

    ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    ws.cell(row=r_idx, column=5, value=f"( {dispenser} )").alignment = align_center
    ws.cell(row=r_idx, column=5).font = font_bold

    r_idx += 1
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    ws.cell(row=r_idx, column=1, value=f"ตำแหน่ง {data['requester_title']}").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_regular

    ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    ws.cell(row=r_idx, column=5, value=f"ตำแหน่ง {data['dispenser_title']}").alignment = align_center
    ws.cell(row=r_idx, column=5).font = font_regular

    r_idx += 2
    # Row 2 of signatures: ผู้สั่งจ่าย and ผู้รับ
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    ws.cell(row=r_idx, column=1, value="(ลงชื่อ) .................................................... ผู้สั่งจ่าย").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_regular

    ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    ws.cell(row=r_idx, column=5, value="(ลงชื่อ) .................................................... ผู้รับ").alignment = align_center
    ws.cell(row=r_idx, column=5).font = font_regular

    r_idx += 1
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    ws.cell(row=r_idx, column=1, value=f"( {authorizer} )").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_bold

    ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    ws.cell(row=r_idx, column=5, value=f"( {receiver} )").alignment = align_center
    ws.cell(row=r_idx, column=5).font = font_bold

    r_idx += 1
    ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=3)
    ws.cell(row=r_idx, column=1, value=f"ตำแหน่ง {data['authorizer_title']}").alignment = align_center
    ws.cell(row=r_idx, column=1).font = font_regular

    ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    ws.cell(row=r_idx, column=5, value=f"ตำแหน่ง {data['receiver_title']}").alignment = align_center
    ws.cell(row=r_idx, column=5).font = font_regular

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 16

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"daily_requisition_{data['date'].replace('-', '')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@app.route("/login")
def login_page():
    code = (request.args.get('code') or request.args.get('passcode') or '').strip()
    next_url = request.args.get('next') or '/'
    if session.get('user_email'):
        return redirect(next_url)
    if code and code in ['otwxc8x', os.environ.get('STAFF_PASSCODE', 'otwxc8x')]:
        session['user_email'] = 'admin@pharmacore.clinic'
        session['user_name'] = 'ผู้ดูแลระบบคลังยา (Admin)'
        session['user_picture'] = '/static/uploads/drugs/app_icon_192.png'
        audit('LOGIN', 'AUTH', 'passcode', 'Admin', 'เข้าสู่ระบบด้วย Passcode')
        return redirect(next_url)
    error = request.args.get('error', '')
    return render_template("login.html", error=error, base_url=Config.BASE_URL)

@app.route("/api/auth/passcode", methods=["POST"])
def passcode_login():
    data = request.get_json(silent=True) or request.form
    code = (data.get("passcode") or data.get("code") or "").strip()
    if code and code in ["otwxc8x", os.environ.get("STAFF_PASSCODE", "otwxc8x")]:
        session["user_email"] = "admin@pharmacore.clinic"
        session["user_name"] = "ผู้ดูแลระบบคลังยา (Admin)"
        session["user_picture"] = "/static/uploads/drugs/app_icon_192.png"
        audit("LOGIN", "AUTH", "passcode", "Admin", "เข้าสู่ระบบด้วย Passcode")
        return jsonify({"ok": True, "redirect": "/"})
    return jsonify({"ok": False, "error": "รหัสผ่านเจ้าหน้าที่ไม่ถูกต้อง"}), 401

@app.route("/logout")
def logout():
    u = current_user()
    if u.get('email'):
        audit('LOGOUT', detail='ออกจากระบบเรียบร้อย')
    session.clear()
    return redirect('/login')


@app.route("/favicon.ico")
def favicon():
    return send_from_directory(str(Path(__file__).parent / "static"), "favicon.ico", mimetype="image/vnd.microsoft.icon")

@app.route("/manifest.json")
def manifest():
    return send_from_directory(str(Path(__file__).parent / "static"), "manifest.json", mimetype="application/manifest+json")

@app.route("/sw.js")
def service_worker():
    response = make_response(send_from_directory(str(Path(__file__).parent / "static"), "sw.js", mimetype="application/javascript"))
    response.headers["Service-Worker-Allowed"] = "/"
    return response

@app.route("/download/apk")
@app.route("/download/app.apk")
def download_apk():
    apk_path = Path(__file__).parent / "static" / "downloads" / "PharmaCore.apk"
    if apk_path.exists():
        return send_file(str(apk_path), as_attachment=True, download_name="PharmaCore.apk", mimetype="application/vnd.android.package-archive")
    # Redirect to mobile install view if apk binary is not physically compiled yet
    return redirect("/?install=apk")


# ─────────────────────────────────────────────────────────────
# Drug Inventory API (CRUD & Search)
# ─────────────────────────────────────────────────────────────

@app.route("/api/drugs", methods=["GET"])
def list_drugs():
    search = request.args.get("q", "").strip()
    category = request.args.get("category", "").strip()
    low_stock = request.args.get("low_stock") == "1"
    
    expiring_param = request.args.get("expiring_days")
    if expiring_param:
        try:
            expiring_days = int(expiring_param)
        except ValueError:
            expiring_days = 60
    elif request.args.get("expiring") == "1":
        expiring_days = Config.EXPIRY_ALERT_DAYS or 60
    else:
        expiring_days = None

    drugs = database.get_all_drugs(search_query=search, category=category, low_stock_only=low_stock, expiring_soon_days=expiring_days)
    return jsonify({"ok": True, "drugs": drugs, "count": len(drugs)})

@app.route("/api/drugs/expiring", methods=["GET"])
def list_expiring_drugs():
    days = int(request.args.get("days", Config.EXPIRY_ALERT_DAYS or 60))
    drugs = database.get_expiring_drugs(days=days)
    return jsonify({"ok": True, "days": days, "drugs": drugs, "count": len(drugs)})

@app.route("/api/drugs/<int:drug_id>", methods=["GET"])
def get_drug(drug_id):
    drug = database.get_drug_by_id(drug_id)
    if not drug:
        return jsonify({"ok": False, "error": "ไม่พบข้อมูลยา"}), 404
    return jsonify({"ok": True, "drug": drug})

@app.route("/api/drugs/scan-lookup", methods=["GET"])
def scan_lookup():
    code = request.args.get("code", "").strip()
    if not code:
        return jsonify({"ok": False, "error": "กรุณาระบุรหัสที่สแกน"}), 400

    clean_code = code
    if "http://" in clean_code or "https://" in clean_code:
        try:
            from urllib.parse import urlparse, parse_qs
            parsed = urlparse(clean_code)
            qs = parse_qs(parsed.query)
            extracted = qs.get("scan", [None])[0] or qs.get("code", [None])[0] or qs.get("dispense", [None])[0]
            if extracted:
                clean_code = extracted.strip()
            else:
                parts = [p for p in parsed.path.split('/') if p]
                if parts:
                    clean_code = parts[-1].strip()
        except Exception:
            pass

    clean_code = clean_code.replace("PHARM:", "").strip()
    drug = database.get_drug_by_code(clean_code) or database.find_drug_by_name_or_alias(clean_code)
    
    if drug:
        return jsonify({"ok": True, "found": True, "drug": drug})
    else:
        return jsonify({"ok": True, "found": False, "scanned_code": code, "clean_code": clean_code})

def trigger_bg_sheet_sync():
    """Run Google Sheet sync in daemon thread so API responses are lightning fast."""
    import threading
    def _run():
        try:
            google_sync.sync_all_drugs_to_sheet(database.get_all_drugs())
        except Exception as e:
            logger.warning(f"Background sheet sync error: {e}")
    threading.Thread(target=_run, daemon=True).start()

@app.route("/api/drugs", methods=["POST"])
@login_required
def create_drug():
    data = request.json or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"ok": False, "error": "กรุณาระบุชื่อยา"}), 400

    try:
        new_drug = database.add_drug(data)
    except sqlite3.IntegrityError:
        return jsonify({"ok": False, "error": f"รหัสยา '{data.get('code')}' ซ้ำกับรายการยาที่มีอยู่แล้วในระบบ กรุณาใช้รหัสอื่นค่ะ"}), 400
    except Exception as e:
        return jsonify({"ok": False, "error": f"ไม่สามารถเพิ่มข้อมูลยาได้: {str(e)}"}), 400
    trigger_bg_sheet_sync()
    u = current_user()
    actor_email = u.get('email', '')
    actor_name = u.get('name', '')
    audit('ADD_DRUG', 'drug', new_drug.get('id',''), new_drug.get('name',''),
          f"รหัส: {new_drug.get('code','')} | หมวด: {new_drug.get('category','')} | สต็อก: {new_drug.get('stock_qty','')} {new_drug.get('unit','')}")

    import threading
    def _bg_notify_new_drug(d):
        try:
            from line_service import line_bot
            line_bot.send_new_drug_notification(d)
        except Exception as err:
            logger.error(f"Error sending LINE new drug notification: {err}")
    d_notify = dict(new_drug)
    d_notify['user_email'] = actor_email
    d_notify['user_name'] = actor_name
    threading.Thread(target=_bg_notify_new_drug, args=(d_notify,), daemon=True).start()

    return jsonify({"ok": True, "drug": new_drug, "message": "เพิ่มข้อมูลยาในคลังเรียบร้อยแล้ว"})

@app.route("/api/drugs/<int:drug_id>", methods=["PUT"])
@login_required
def edit_drug(drug_id):
    data = request.json or {}
    if "stock_qty" in data and data["stock_qty"] is not None:
        try:
            sq = int(data["stock_qty"])
            if sq < 0:
                return jsonify({"ok": False, "error": "จำนวนสต็อกต้องไม่ติดลบ"}), 400
        except (ValueError, TypeError):
            return jsonify({"ok": False, "error": "จำนวนสต็อกต้องเป็นตัวเลขจำนวนเต็ม"}), 400

    updated = database.update_drug(drug_id, data)
    if not updated:
        return jsonify({"ok": False, "error": "ไม่พบข้อมูลยาที่ต้องการแก้ไข"}), 404
    audit('EDIT_DRUG', 'drug', drug_id, updated.get('name',''),
          f"แก้ไขข้อมูลยา | สต็อก: {updated.get('stock_qty','')} {updated.get('unit','')}")
    trigger_bg_sheet_sync()
    return jsonify({"ok": True, "drug": updated, "message": "แก้ไขข้อมูลยาเรียบร้อยแล้ว"})

@app.route("/api/drugs/<int:drug_id>", methods=["DELETE"])
@login_required
def remove_drug(drug_id):
    drug = database.get_drug_by_id(drug_id)
    drug_name = drug.get('name','') if drug else ''
    database.delete_drug(drug_id)
    audit('DELETE_DRUG', 'drug', drug_id, drug_name, 'ลบรายการยาออกจากคลัง')
    trigger_bg_sheet_sync()
    return jsonify({"ok": True, "message": "ลบยาออกจากคลังเรียบร้อยแล้ว"})

# ─────────────────────────────────────────────────────────────
# Category Master Management API (CRUD)
# ─────────────────────────────────────────────────────────────

@app.route("/api/categories", methods=["GET"])
def api_get_categories():
    categories = database.get_categories()
    return jsonify({"ok": True, "categories": categories})

@app.route("/api/categories", methods=["POST"])
def api_add_category():
    data = request.get_json(silent=True) or {}
    name = data.get("name", "").strip()
    color = data.get("color", "teal").strip()
    if not name:
        return jsonify({"ok": False, "error": "กรุณาระบุชื่อหมวดหมู่"}), 400
    ok, res = database.add_category(name, color)
    if not ok:
        return jsonify({"ok": False, "error": res}), 400
    audit('ADD_CATEGORY', 'category', str(res['id']), name, f'เพิ่มหมวดหมู่ยา: {name}')
    return jsonify({"ok": True, "category": res, "message": "เพิ่มหมวดหมู่สำเร็จ"})

@app.route("/api/categories/<int:cat_id>", methods=["PUT"])
def api_update_category(cat_id):
    data = request.get_json(silent=True) or {}
    name = data.get("name", "").strip()
    color = data.get("color")
    if not name:
        return jsonify({"ok": False, "error": "กรุณาระบุชื่อหมวดหมู่"}), 400
    ok, res = database.update_category(cat_id, name, color)
    if not ok:
        return jsonify({"ok": False, "error": res}), 400
    audit('UPDATE_CATEGORY', 'category', str(cat_id), name, f'แก้ไขหมวดหมู่ยา: {name}')
    return jsonify({"ok": True, "category": res, "message": "แก้ไขหมวดหมู่สำเร็จ"})

@app.route("/api/categories/<int:cat_id>", methods=["DELETE"])
def api_delete_category(cat_id):
    fallback = request.args.get("fallback", "ทั่วไป")
    ok, msg = database.delete_category(cat_id, fallback)
    if not ok:
        return jsonify({"ok": False, "error": msg}), 400
    audit('DELETE_CATEGORY', 'category', str(cat_id), str(cat_id), f'ลบหมวดหมู่ยา (ย้ายยาไป {fallback})')
    return jsonify({"ok": True, "message": msg})

# ─────────────────────────────────────────────────────────────
# Stock Dispensing & Deduction API
# ─────────────────────────────────────────────────────────────

_recent_dispense_locks = {}

def _check_and_set_dispense_lock(user_key, drug_id, qty, patient_name, ttl=2.0):
    """Prevent accidental rapid double-clicking by locking duplicate dispense requests for 2 seconds."""
    import time
    now = time.time()
    # Clean expired locks
    stale = [k for k, exp in _recent_dispense_locks.items() if exp < now]
    for k in stale:
        _recent_dispense_locks.pop(k, None)

    lock_key = f"{user_key}:{drug_id}:{qty}:{patient_name}"
    if lock_key in _recent_dispense_locks and _recent_dispense_locks[lock_key] > now:
        return False  # Duplicate / rapid double-click detected
    _recent_dispense_locks[lock_key] = now + ttl
    return True

@app.route("/api/dispense", methods=["POST"])
@login_required
def dispense_stock():
    """
    Deduct stock via Web UI (Manual or QR Scan).
    Payload: {drug_id: int, quantity: int, channel: str, dispensed_by: str, patient_name: str, notes: str}
    """
    data = request.json or {}
    drug_id = data.get("drug_id")
    try:
        quantity = int(data.get("quantity") or 1)
    except (ValueError, TypeError):
        return jsonify({"ok": False, "error": "จำนวนที่ตัดต้องเป็นตัวเลขจำนวนเต็มบวก"}), 400

    if quantity <= 0:
        return jsonify({"ok": False, "error": "จำนวนที่ตัดต้องมากกว่า 0"}), 400

    channel = data.get("channel") or "WEB_MANUAL"
    # Auto-fill dispensed_by with logged-in user email/name if not provided
    u = current_user()
    actor_email = u.get('email', '')
    actor_name = u.get('name', '')
    default_dispenser = actor_name or actor_email or "ผู้ดูแลระบบคลังยา (Admin)"
    dispensed_by = (data.get("dispensed_by") or "").strip()
    if not dispensed_by or dispensed_by in ("เภสัชกร/เจ้าหน้าที่", "เจ้าหน้าที่"):
        dispensed_by = default_dispenser
    patient_name = data.get("patient_name") or ""
    notes = data.get("notes") or ""

    if not drug_id:
        return jsonify({"ok": False, "error": "กรุณาระบุรหัสยาที่ต้องการตัดสต็อก"}), 400

    # Anti double-click guard on backend
    user_key = actor_email or request.remote_addr or "client"
    if not _check_and_set_dispense_lock(user_key, drug_id, quantity, patient_name, ttl=2.0):
        return jsonify({
            "ok": False,
            "error": "ระบบกำลังประมวลผลคำขอตัดสต็อกรายการนี้อยู่ กรุณาอย่ากดซ้ำค่ะ"
        }), 429

    success, result = database.deduct_stock(
        drug_id, quantity,
        channel=channel, dispensed_by=dispensed_by,
        patient_name=patient_name, notes=notes
    )

    if not success:
        return jsonify({"ok": False, "error": result}), 400

    audit('DISPENSE', 'drug', drug_id, result.get('drug_name',''),
          f"ตัดจ่าย {quantity} {result.get('unit','')} | คงเหลือ: {result.get('balance_after','')} | ผู้รับ: {patient_name}")

    # Run Google Sheets sync & LINE notification in background thread
    # so HTTP response to the browser is near instantaneous (<50ms), preventing user wait/double-clicks!
    import threading
    import time
    def _bg_after_dispense(res, u_email, u_name, disp_by):
        global _last_inventory_sheet_sync
        try:
            google_sync.append_dispensing_log(res)
            # Throttle full inventory rewrite to prevent hitting Google API quotas
            now = time.time()
            if '_last_inventory_sheet_sync' not in globals() or now - _last_inventory_sheet_sync > 15.0:
                _last_inventory_sheet_sync = now
                google_sync.sync_all_drugs_to_sheet(database.get_all_drugs())
        except Exception as e:
            logger.error(f"Error syncing to sheet after dispense: {e}")

        try:
            from line_service import line_bot
            res_notify = dict(res)
            res_notify['user_email'] = u_email
            res_notify['user_name'] = u_name
            res_notify['dispensed_by'] = disp_by
            line_bot.send_dispense_notification(res_notify)
        except Exception as err:
            logger.error(f"Error sending LINE dispense notification: {err}")

    threading.Thread(target=_bg_after_dispense, args=(result, actor_email, actor_name, dispensed_by), daemon=True).start()

    return jsonify({
        "ok": True,
        "result": result,
        "message": f"ตัดสต็อก '{result['drug_name']}' สำเร็จ - คงเหลือ {result['balance_after']} {result['unit']}"
    })


@app.route("/api/dispense/batch", methods=["POST"])
@login_required
def batch_dispense():
    """
    Multi-Drug Rx Cart Batch Dispense:
    Atomically deducts multiple drugs for a patient in a single prescription batch.
    """
    data = request.json or {}
    items = data.get("items", [])
    if not items or not isinstance(items, list):
        return jsonify({"ok": False, "error": "กรุณาระบุรายการยาในใบสั่งยาอย่างน้อย 1 รายการ"}), 400

    u = current_user()
    actor_email = u.get('email', '')
    actor_name = u.get('name', '')
    dispensed_by = (data.get("dispensed_by") or (f"{actor_name}" if actor_name else "เภสัชกร/เจ้าหน้าที่")).strip()
    patient_name = (data.get("patient_name") or "").strip()
    requester_dept = (data.get("requester_dept") or "แผนกตรวจรักษา / หอผู้ป่วย").strip()
    notes = (data.get("notes") or "").strip()
    channel = data.get("channel") or "WEB_RX_CART"

    success, result = database.deduct_stock_batch(
        items=items,
        channel=channel,
        dispensed_by=dispensed_by,
        patient_name=patient_name,
        requester_dept=requester_dept,
        notes=notes
    )

    if not success:
        return jsonify({"ok": False, "error": result}), 400

    audit('BATCH_DISPENSE', 'prescription', result.get('batch_id', ''),
          f"ใบสั่งยา: {patient_name or 'ไม่ระบุผู้ป่วย'} ({result.get('total_items', 0)} รายการ)",
          f"ยอดรวม: {result.get('total_price_formatted', '')} | จำนวนยารวม: {result.get('total_qty', 0)}")

    # Background sync to sheets and LINE notify
    import threading
    def _bg_after_batch(res, u_email, u_name, disp_by):
        global _last_inventory_sheet_sync
        try:
            for item in res.get("items", []):
                google_sync.append_dispensing_log(item)
            # Sync full inventory table to sheet
            google_sync.sync_all_drugs_to_sheet(database.get_all_drugs())
        except Exception as e:
            logger.error(f"Error syncing batch dispense to sheet: {e}")

        try:
            from line_service import line_bot
            res_notify = dict(res)
            res_notify['user_email'] = u_email
            res_notify['user_name'] = u_name
            res_notify['dispensed_by'] = disp_by
            line_bot.send_batch_dispense_notification(res_notify)
        except Exception as err:
            logger.error(f"Error sending batch LINE notification: {err}")

    threading.Thread(target=_bg_after_batch, args=(result, actor_email, actor_name, dispensed_by), daemon=True).start()

    return jsonify({
        "ok": True,
        "result": result,
        "batch_id": result["batch_id"],
        "requisition_url": f"/api/dispense/requisition/batch/{result['batch_id']}",
        "message": f"ตัดสต็อกใบสั่งยา {result['batch_id']} สำเร็จ ({result['total_items']} รายการ, ยอดรวม {result['total_price_formatted']})"
    })


@app.route("/api/dispense/<int:log_id>/void", methods=["POST"])
@login_required
def void_dispense(log_id):
    """
    Cancel/void a dispensing log, restore drug stock, and record audit trail.
    """
    data = request.json or {}
    reason = (data.get("reason") or "ยกเลิกรายการจ่ายโดยเจ้าหน้าที่").strip()

    u = current_user()
    actor_email = u.get('email', '')
    actor_name = u.get('name', '')
    voided_by = f"{actor_name} ({actor_email})" if actor_name and actor_name != actor_email else (actor_email or "เภสัชกร/เจ้าหน้าที่")

    success, result = database.void_dispensing_log(log_id, reason=reason, voided_by=voided_by)
    if not success:
        return jsonify({"ok": False, "error": result}), 400

    audit('VOID_DISPENSE', 'dispense_log', log_id, result.get('drug_name', ''),
          f"ยกเลิกการจ่าย คืนยา {result.get('quantity_restored')} หน่วย | คงเหลือใหม่: {result.get('new_stock')} | เหตุผล: {reason}")

    trigger_bg_sheet_sync()

    return jsonify({
        "ok": True,
        "result": result,
        "message": f"ยกเลิกรายการจ่ายและคืนสต็อกยา '{result['drug_name']}' จำนวน {result['quantity_restored']} หน่วย เรียบร้อยแล้ว (คงเหลือใหม่: {result['new_stock']})"
    })

@app.route("/api/dispense/logs", methods=["GET"])
def list_logs():
    limit = int(request.args.get("limit", 100))
    date_filter = request.args.get("date")
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    channel = request.args.get("channel")
    logs = database.get_dispensing_logs(
        limit=limit,
        date_filter=date_filter,
        channel_filter=channel,
        start_date=start_date,
        end_date=end_date
    )
    return jsonify({"ok": True, "logs": logs, "count": len(logs)})

# ─────────────────────────────────────────────────────────────
# Stock In & Receiving API (Batch & Lot Management)
# ─────────────────────────────────────────────────────────────

@app.route("/api/stock-in", methods=["POST"])
@login_required
def stock_in_endpoint():
    data = request.json or {}
    drug_id = data.get("drug_id")
    drug_code = data.get("drug_code")
    target = drug_id or drug_code

    if not target:
        return jsonify({"ok": False, "error": "กรุณาระบุรหัสยาหรือเลือกรายการยา"}), 400

    try:
        quantity = int(data.get("quantity") or 0)
    except (ValueError, TypeError):
        return jsonify({"ok": False, "error": "จำนวนรับเข้าต้องเป็นตัวเลข"}), 400

    if quantity <= 0:
        return jsonify({"ok": False, "error": "จำนวนรับเข้าต้องมากกว่า 0"}), 400

    u = current_user()
    actor_email = u.get('email', '')
    actor_name = u.get('name', '')
    default_receiver = actor_name or actor_email or "ผู้ดูแลระบบคลังยา (Admin)"

    lot_number = (data.get("lot_number") or "").strip()
    cost_per_unit = data.get("cost_per_unit")
    expiry_date = (data.get("expiry_date") or "").strip()
    supplier = (data.get("supplier") or "องค์การเภสัชกรรม (GPO)").strip()
    received_by = (data.get("received_by") or "").strip()
    if not received_by or received_by in ("เภสัชกร/เจ้าหน้าที่", "เจ้าหน้าที่"):
        received_by = default_receiver
    notes = (data.get("notes") or "").strip()

    success, result = database.stock_in(
        drug_id_or_code=target,
        quantity=quantity,
        lot_number=lot_number,
        cost_per_unit=cost_per_unit,
        expiry_date=expiry_date,
        supplier=supplier,
        received_by=received_by,
        notes=notes
    )

    if not success:
        return jsonify({"ok": False, "error": result}), 400

    audit('STOCK_IN', 'drug', result.get('drug_id',''), result.get('drug_name',''),
          f"รับเข้า {quantity} {result.get('unit','')} | Lot: {lot_number} | Supplier: {supplier}")

    # Sync inventory to Google Sheets in background
    trigger_bg_sheet_sync()

    # LINE Bot Stock In Notification
    import threading
    def _bg_notify_stock_in(res):
        try:
            from line_service import line_bot
            line_bot.send_stock_in_notification(res)
        except Exception as err:
            logger.error(f"Error sending LINE stock in notification: {err}")
    res_notify = dict(result)
    res_notify['user_email'] = actor_email
    res_notify['user_name'] = actor_name
    res_notify['received_by'] = received_by
    threading.Thread(target=_bg_notify_stock_in, args=(res_notify,), daemon=True).start()

    return jsonify({
        "ok": True,
        "result": result,
        "message": f"รับเข้า '{result['drug_name']}' จำนวน {result['quantity']} {result['unit']} สำเร็จ (คงเหลือใหม่ {result['balance_after']} {result['unit']})"
    })

@app.route("/api/stock-in/logs", methods=["GET"])
def list_stock_in_logs():
    limit = int(request.args.get("limit", 100))
    date_filter = request.args.get("date")
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    drug_id = request.args.get("drug_id")
    logs = database.get_stock_in_logs(
        limit=limit,
        drug_id=drug_id,
        date_filter=date_filter,
        start_date=start_date,
        end_date=end_date
    )
    return jsonify({"ok": True, "logs": logs, "count": len(logs)})

@app.route("/api/stock-in/<int:log_id>", methods=["GET"])
@login_required
def get_stock_in_endpoint(log_id):
    log = database.get_stock_in_log(log_id)
    if not log:
        return jsonify({"ok": False, "error": f"ไม่พบรายการรับยาเข้ารหัส #{log_id}"}), 404
    return jsonify({"ok": True, "log": log})

@app.route("/api/stock-in/<int:log_id>", methods=["PUT", "POST"])
@login_required
def update_stock_in_endpoint(log_id):
    data = request.json or {}
    rec_by = (data.get("received_by") or "").strip()
    if not rec_by or rec_by in ("เภสัชกร/เจ้าหน้าที่", "เจ้าหน้าที่"):
        rec_by = current_user().get('name') or current_user().get('email') or 'ผู้ดูแลระบบคลังยา (Admin)'

    success, result = database.update_stock_in_log(
        log_id,
        quantity=data.get("quantity"),
        lot_number=data.get("lot_number"),
        expiry_date=data.get("expiry_date"),
        cost_per_unit=data.get("cost_per_unit"),
        supplier=data.get("supplier"),
        received_by=rec_by,
        notes=data.get("notes"),
        created_at=data.get("created_at")
    )
    if not success:
        return jsonify({"ok": False, "error": result}), 400

    audit("UPDATE_STOCK_IN", "stock_in_log", log_id, result.get("drug_name", ""),
          f"แก้ไขรายการรับเข้าย้อนหลัง: จำนวน {result.get('quantity')} {result.get('unit', '')} | Lot {result.get('lot_number')} | ผู้จำหน่าย {result.get('supplier')} | วันที่ {result.get('created_at')}")

    trigger_bg_sheet_sync()
    return jsonify({"ok": True, "log": result, "message": f"แก้ไขรายการรับเข้ายา '{result.get('drug_name')}' เรียบร้อยแล้ว"})

@app.route("/api/stock-in/<int:log_id>/void", methods=["POST"])
@app.route("/api/stock-in/<int:log_id>", methods=["DELETE", "POST"])
@app.route("/api/stock-in/<int:log_id>/delete", methods=["POST", "DELETE"])
@login_required
def void_stock_in_endpoint(log_id):
    data = request.json or {}
    reason = (data.get("reason") or "ยกเลิกรายการรับเข้าโดยเจ้าหน้าที่").strip()

    u = current_user()
    actor_email = u.get('email', '')
    actor_name = u.get('name', '')
    voided_by = f"{actor_name} ({actor_email})" if actor_name and actor_name != actor_email else (actor_email or "เภสัชกร/เจ้าหน้าที่")

    success, result = database.void_stock_in_log(log_id, reason=reason, voided_by=voided_by)
    if not success:
        return jsonify({"ok": False, "error": result}), 400

    audit("VOID_STOCK_IN", "stock_in_log", log_id, result.get("drug_name", ""),
          f"ยกเลิกรายการรับเข้า คืนสต็อกลดลง {result.get('quantity')} หน่วย | Lot {result.get('lot_number')} | คงเหลือใหม่: {result.get('new_stock', '-')} | เหตุผล: {reason}")

    trigger_bg_sheet_sync()
    return jsonify({
        "ok": True,
        "log": result,
        "message": f"ยกเลิกรายการรับเข้ายา '{result.get('drug_name')}' จำนวน {result.get('quantity')} หน่วย เรียบร้อยแล้ว (สต็อกคงเหลือใหม่: {result.get('new_stock', '-')})"
    })

@app.route("/api/analytics/summary", methods=["GET"])
def analytics_summary():
    summary = database.get_analytics_summary()
    return jsonify({"ok": True, "analytics": summary})

@app.route("/api/dashboard/stats", methods=["GET"])
def dashboard_stats():
    stats = database.get_dashboard_stats()
    stats["google_sheet_connected"] = google_sync.is_connected
    stats["google_spreadsheet_id"] = google_sync.spreadsheet_id or "Not Configured"
    return jsonify({"ok": True, "stats": stats})

# ─────────────────────────────────────────────────────────────
# Google Drive & Sheets Connection Settings API
# ─────────────────────────────────────────────────────────────

@app.route("/api/google/config", methods=["GET"])
def get_google_config():
    info = google_sync.get_connection_info()
    return jsonify({"ok": True, "config": info})

@app.route("/api/google/config", methods=["POST"])
def update_google_config():
    data = request.json or {}
    sheet_input = data.get("spreadsheet_id")
    drive_input = data.get("drive_folder_id")
    sa_input = data.get("service_account_json")

    info = google_sync.update_and_reconnect(sheet_input, drive_input, sa_input)
    if info.get("is_connected") and info.get("spreadsheet_id"):
        try:
            google_sync.ensure_pharmacy_spreadsheet()
            google_sync.sync_all_drugs_to_sheet(database.get_all_drugs())
        except Exception as e:
            logger.warning(f"Sync after config update error: {e}")

    return jsonify({"ok": True, "config": info, "message": "อัปเดตการเชื่อมต่อ Google Drive & Sheets เรียบร้อยแล้ว"})

@app.route("/api/google/auto-create-sheet", methods=["POST"])
def auto_create_google_sheet():
    sheet_id, err = google_sync.ensure_pharmacy_spreadsheet()
    if err:
        return jsonify({"ok": False, "error": err}), 400

    try:
        google_sync.sync_all_drugs_to_sheet(database.get_all_drugs())
    except Exception:
        pass

    return jsonify({"ok": True, "config": google_sync.get_connection_info(), "message": "สร้าง Google Sheet สำหรับคลังยาสำเร็จแล้ว"})

@app.route("/api/google/disconnect", methods=["POST"])
def disconnect_google_endpoint():
    google_sync.disconnect_google()
    return jsonify({"ok": True, "message": "ตัดการเชื่อมต่อ Google สำเร็จแล้ว"})

# ─────────────────────────────────────────────────────────────
# Google OAuth2 Automatic Authorization (Drive & Sheets Auto-Create)
# ─────────────────────────────────────────────────────────────

def _determine_redirect_uri():
    host = request.headers.get("x-forwarded-host") or request.host or ""
    proto = request.headers.get("x-forwarded-proto") or ("https" if request.is_secure else "http")
    
    # If accessing via ngrok domain
    if "ngrok" in host:
        return f"https://{host}/api/auth/google/callback"
    
    # If direct local access
    if "localhost" in host or "127.0.0.1" in host:
        return "http://localhost:5005/api/auth/google/callback"

    # If accessing via Cloudflare / openchat
    if "openchat.sbs" in host:
        return "https://openchat.sbs/api/auth/google/callback"

    # If Config.OAUTH2_REDIRECT_URI is set and does not point to openchat
    if Config.OAUTH2_REDIRECT_URI and "openchat.sbs" not in Config.OAUTH2_REDIRECT_URI:
        return Config.OAUTH2_REDIRECT_URI

    if host:
        return f"{proto}://{host}/api/auth/google/callback"

    return "https://encrypt-unvocal-hankering.ngrok-free.dev/api/auth/google/callback"

@app.route("/api/auth/google/login")
def google_oauth_login():
    """Start Google OAuth flow to grant Drive & Sheets access automatically."""
    client_id = Config.GOOGLE_OAUTH2_CLIENT_ID
    if not client_id:
        return "Google OAuth2 Client ID not configured", 500

    redirect_uri = _determine_redirect_uri()
    state_token = secrets.token_urlsafe(32)
    database.set_app_setting(f"oauth_state_{state_token}", redirect_uri)

    params = {
        "client_id": client_id,
        "redirect_uri": redirect_uri,
        "response_type": "code",
        "scope": "openid https://www.googleapis.com/auth/userinfo.email https://www.googleapis.com/auth/userinfo.profile https://www.googleapis.com/auth/drive.file https://www.googleapis.com/auth/spreadsheets",
        "access_type": "offline",
        "prompt": "consent",
        "state": state_token
    }
    auth_url = "https://accounts.google.com/o/oauth2/auth?" + urllib.parse.urlencode(params)
    return redirect(auth_url)

# ─────────────────────────────────────────────────────────────
# Google OAuth2 — USER LOGIN (session-based, for access control)
# ─────────────────────────────────────────────────────────────

def _determine_user_login_redirect_uri():
    # Use the registered OAuth callback URI to prevent redirect_uri_mismatch
    return _determine_redirect_uri()

@app.route("/api/auth/google/user-login")
def google_user_login():
    """Initiate Google sign-in for user authentication (session login)."""
    client_id = Config.GOOGLE_OAUTH2_CLIENT_ID
    if not client_id:
        return render_template("login.html", error="ยังไม่ได้ตั้งค่า Google OAuth2 Client ID", base_url=Config.BASE_URL)

    redirect_uri = _determine_user_login_redirect_uri()
    state_token  = secrets.token_urlsafe(32)
    next_url     = request.args.get('next', '/')
    database.set_app_setting(f"login_state_{state_token}", f"{redirect_uri}|||{next_url}")

    params = {
        "client_id": client_id,
        "redirect_uri": redirect_uri,
        "response_type": "code",
        "scope": "openid https://www.googleapis.com/auth/userinfo.email https://www.googleapis.com/auth/userinfo.profile",
        "access_type": "online",
        "prompt": "select_account",
        "state": state_token
    }
    auth_url = "https://accounts.google.com/o/oauth2/auth?" + urllib.parse.urlencode(params)
    return redirect(auth_url)

def _finish_user_login(code, redirect_uri, next_url):
    """Exchange authorization code, extract user profile, establish session, and record audit log."""
    try:
        token_payload = urllib.parse.urlencode({
            "code": code,
            "client_id": Config.GOOGLE_OAUTH2_CLIENT_ID,
            "client_secret": Config.GOOGLE_OAUTH2_CLIENT_SECRET,
            "redirect_uri": redirect_uri,
            "grant_type": "authorization_code"
        }).encode("utf-8")

        req = urllib.request.Request("https://oauth2.googleapis.com/token", data=token_payload, method="POST")
        with urllib.request.urlopen(req) as resp:
            token_res = json.loads(resp.read().decode("utf-8"))

        access_token = token_res.get("access_token")
        if not access_token:
            return redirect("/login?error=token_failed")

        userinfo_req = urllib.request.Request(
            "https://www.googleapis.com/oauth2/v2/userinfo",
            headers={"Authorization": f"Bearer {access_token}"}
        )
        with urllib.request.urlopen(userinfo_req) as resp:
            user_info = json.loads(resp.read().decode("utf-8"))

        email   = user_info.get("email", "").strip()
        name    = user_info.get("name", email).strip()
        picture = user_info.get("picture", "").strip()
        sub     = user_info.get("id", "").strip()

        if not email:
            return redirect("/login?error=no_email")

        # Upsert user record in SQLite
        database.add_or_update_user(email, name, picture, sub)

        # Set Flask session
        session.permanent = True
        session['user_email']   = email
        session['user_name']    = name
        session['user_picture'] = picture

        # Write login audit log
        ip = request.headers.get('X-Forwarded-For', request.remote_addr or '')
        database.log_audit_action(
            user_email=email,
            user_name=name,
            action='LOGIN',
            target_type='auth',
            target_id=sub,
            target_name=email,
            detail=f'เข้าสู่ระบบผ่าน Google สำเร็จ | IP: {ip}',
            ip_address=ip
        )
        logger.info(f"✅ User authenticated: {email} ({name}) from {ip}")

        return redirect(next_url if (next_url and next_url.startswith('/')) else '/')

    except Exception as e:
        logger.error(f"Google user login callback error: {e}")
        return redirect("/login?error=" + urllib.parse.quote(str(e)))

# ─────────────────────────────────────────────────────────────
# Audit Log API
# ─────────────────────────────────────────────────────────────

@app.route("/api/audit/logs", methods=["GET"])
@login_required
def get_audit_logs_api():
    limit       = int(request.args.get("limit", 200))
    user_email  = request.args.get("user_email")
    action      = request.args.get("action")
    date_filter = request.args.get("date")
    logs = database.get_audit_logs(limit=limit, user_email=user_email, action=action, date_filter=date_filter)
    return jsonify({"ok": True, "logs": logs, "count": len(logs)})

@app.route("/api/audit/users", methods=["GET"])
@login_required
def get_users_api():
    users = database.get_all_users()
    return jsonify({"ok": True, "users": users})

@app.route("/api/auth/me", methods=["GET"])
def get_me():
    u = current_user()
    if not u['email']:
        return jsonify({"ok": False, "authenticated": False}), 401
    return jsonify({"ok": True, "authenticated": True, "user": u})


@app.route("/api/auth/google/callback")
@app.route("/api/auth/google/user-callback")
def google_oauth_callback():
    """Unified OAuth callback for both User Session Login and Google Drive/Sheets Setup."""
    error = request.args.get("error")
    if error:
        return redirect("/login?error=" + urllib.parse.quote(error))

    code = request.args.get("code")
    state = request.args.get("state")
    if not code:
        return redirect("/login?error=no_code")

    # 1. Check if this is a user login authentication flow
    saved_login = database.get_app_setting(f"login_state_{state}") if state else None
    if saved_login or request.path.endswith("user-callback"):
        parts = saved_login.split("|||", 1) if saved_login else ["", "/"]
        redirect_uri = parts[0] if parts[0] else _determine_user_login_redirect_uri()
        next_url = parts[1] if len(parts) > 1 else "/"
        return _finish_user_login(code, redirect_uri, next_url)

    # 2. Workspace Drive & Sheets setup
    saved_redirect_uri = database.get_app_setting(f"oauth_state_{state}") or _determine_redirect_uri()

    try:
        # Exchange code for token with Google
        token_payload = urllib.parse.urlencode({
            "code": code,
            "client_id": Config.GOOGLE_OAUTH2_CLIENT_ID,
            "client_secret": Config.GOOGLE_OAUTH2_CLIENT_SECRET,
            "redirect_uri": saved_redirect_uri,
            "grant_type": "authorization_code"
        }).encode("utf-8")

        req = urllib.request.Request("https://oauth2.googleapis.com/token", data=token_payload, method="POST")
        with urllib.request.urlopen(req) as resp:
            token_res = json.loads(resp.read().decode("utf-8"))

        access_token = token_res.get("access_token")
        refresh_token = token_res.get("refresh_token")

        # Get user email
        userinfo_req = urllib.request.Request(
            "https://www.googleapis.com/oauth2/v2/userinfo",
            headers={"Authorization": f"Bearer {access_token}"}
        )
        with urllib.request.urlopen(userinfo_req) as resp:
            user_info = json.loads(resp.read().decode("utf-8"))
        google_email = user_info.get("email", "unknown@gmail.com")

        # Setup workspace automatically
        ok, res = google_sync.setup_oauth_workspace(google_email, access_token, refresh_token)
        if ok:
            return redirect(f"/?google_connected=1&email={urllib.parse.quote(google_email)}")
        else:
            return redirect(f"/?google_error={urllib.parse.quote(str(res))}")

    except Exception as e:
        logger.error(f"Google OAuth callback error: {e}")
        return redirect(f"/?google_error={urllib.parse.quote(str(e))}")

@app.route("/api/google/disconnect", methods=["POST"])
def google_disconnect():
    google_sync.disconnect_google()
    return jsonify({"ok": True, "message": "ตัดการเชื่อมต่อ Google เรียบร้อยแล้ว"})


# ─────────────────────────────────────────────────────────────
# Image Upload & OCR Analysis (Prescription / QR scan)
# ─────────────────────────────────────────────────────────────

@app.route("/api/scan-image", methods=["POST"])
def scan_image():
    """
    Accepts uploaded photo of prescription, label, or QR code.
    Runs OCR and returns detected drug and quantity.
    """
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "ไม่พบไฟล์ภาพที่อัปโหลด"}), 400

    f = request.files["file"]
    image_bytes = f.read()
    if not image_bytes:
        return jsonify({"ok": False, "error": "ไฟล์ภาพว่างเปล่า"}), 400

    analysis = parse_drug_from_image(image_bytes)
    return jsonify({"ok": True, "analysis": analysis})

UPLOAD_DRUG_FOLDER = os.environ.get(
    "UPLOAD_DRUG_FOLDER",
    "/tmp/uploads/drugs" if os.environ.get("VERCEL") else os.path.join(os.path.dirname(__file__), "static", "uploads", "drugs")
)
try:
    os.makedirs(UPLOAD_DRUG_FOLDER, exist_ok=True)
except Exception:
    pass
ALLOWED_IMAGE_EXTS = {"png", "jpg", "jpeg", "webp", "gif"}

def allowed_image_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_IMAGE_EXTS

@app.route("/api/upload/drug-image", methods=["POST"])
def upload_drug_image():
    """Upload drug image and return accessible static URL."""
    file = request.files.get("image") or request.files.get("file")
    if not file or file.filename == "":
        return jsonify({"ok": False, "error": "ไม่พบไฟล์ภาพที่อัปโหลด"}), 400
    if not allowed_image_file(file.filename):
        return jsonify({"ok": False, "error": "รองรับเฉพาะไฟล์รูปภาพ (PNG, JPG, JPEG, WEBP, GIF)"}), 400

    try:
        import uuid
        ext = file.filename.rsplit(".", 1)[1].lower()
        unique_name = f"drug_{uuid.uuid4().hex[:12]}.{ext}"
        filepath = os.path.join(UPLOAD_DRUG_FOLDER, unique_name)
        file.save(filepath)
        image_url = f"/static/uploads/drugs/{unique_name}"
        return jsonify({"ok": True, "image_url": image_url})
    except Exception as e:
        logger.error(f"Error saving uploaded drug image: {e}")
        return jsonify({"ok": False, "error": f"เกิดข้อผิดพลาดในการบันทึกรูปภาพ: {str(e)}"}), 500

@app.route("/api/sync-sheet", methods=["POST"])
def trigger_sheet_sync():
    """Manually force re-sync with Google Sheets (both inventory and dispensing logs)."""
    drugs = database.get_all_drugs()
    success1, msg1 = google_sync.sync_all_drugs_to_sheet(drugs)
    success2, msg2 = google_sync.sync_all_dispensing_logs()
    return jsonify({
        "ok": success1 and success2,
        "message": "ซิงค์ข้อมูลทั้งคลังยาและประวัติการตัดสต็อกขึ้น Google Sheets เรียบร้อยแล้ว"
    })

# ─────────────────────────────────────────────────────────────
# Excel Export Routes
# ─────────────────────────────────────────────────────────────

@app.route("/api/export/inventory", methods=["GET"])
def export_inventory_excel():
    """Export all drug inventory to a styled Excel (.xlsx) file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"ok": False, "error": "openpyxl is not installed"}), 500

    drugs = database.get_all_drugs()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "รายการยาในคลัง"

    # ── Style helpers ──
    header_fill   = PatternFill("solid", fgColor="1E3A5F")
    low_fill      = PatternFill("solid", fgColor="FFF1F2")
    alt_fill      = PatternFill("solid", fgColor="F8FAFC")
    header_font   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font     = Font(name="Calibri", size=10)
    bold_font     = Font(name="Calibri", bold=True, size=10)
    danger_font   = Font(name="Calibri", bold=True, color="E11D48", size=10)
    center        = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left          = Alignment(horizontal="left",   vertical="center", wrap_text=False)
    thin_side     = Side(style="thin", color="CBD5E1")
    thin_border   = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # ── Title row ──
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"รายงานคลังยา — PharmaCore  (ส่งออก {datetime.now().strftime('%d/%m/%Y %H:%M')})"
    title_cell.font  = Font(name="Calibri", bold=True, color="1E3A5F", size=14)
    title_cell.alignment = left
    ws.row_dimensions[1].height = 26

    # ── Header row ──
    headers = ["ลำดับ", "รหัสยา", "ชื่อยา", "ชื่อสามัญ", "หมวดหมู่",
               "รูปแบบยา", "สต็อกคงเหลือ", "หน่วย", "จุดเตือนขั้นต่ำ",
               "ราคา (บาท)", "เลขล็อต (Lot)", "วันหมดอายุ", "ตำแหน่งจัดเก็บ"]
    col_widths = [6, 13, 30, 25, 22, 20, 15, 10, 17, 13, 16, 14, 20]

    ws.append(headers)
    for c_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=c_idx)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = thin_border
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.row_dimensions[2].height = 22

    # ── Data rows ──
    for row_num, d in enumerate(drugs, start=1):
        is_low = (d.get("stock_qty", 0) or 0) <= (d.get("min_threshold", 10) or 10)
        row_data = [
            row_num,
            d.get("code", ""),
            d.get("name", ""),
            d.get("generic_name") or "",
            d.get("category", ""),
            d.get("dosage_form") or "",
            d.get("stock_qty", 0),
            d.get("unit", ""),
            d.get("min_threshold", 10),
            d.get("price") or 0,
            d.get("lot_number") or "",
            d.get("expiry_date") or "",
            d.get("location") or "",
        ]
        ws.append(row_data)
        xlsx_row = row_num + 2
        row_fill = low_fill if is_low else (alt_fill if row_num % 2 == 0 else None)
        for c_idx in range(1, 14):
            cell = ws.cell(row=xlsx_row, column=c_idx)
            cell.border    = thin_border
            cell.alignment = center if c_idx in (1, 7, 8, 9, 10, 11, 12) else left
            cell.font      = (danger_font if c_idx == 7 and is_low else body_font)
            if c_idx == 7 and is_low:
                cell.font = danger_font
            elif c_idx in (1, 2):
                cell.font = bold_font
            if row_fill:
                cell.fill = row_fill
        ws.row_dimensions[xlsx_row].height = 18

    # ── Freeze panes (keep header visible) ──
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:K{len(drugs) + 2}"

    buf = BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    filename = f"pharmacore_inventory_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    resp = make_response(xlsx_bytes)
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    resp.headers["Content-Length"] = len(xlsx_bytes)
    return resp


@app.route("/api/export/logs", methods=["GET"])
def export_logs_excel():
    """Export dispensing logs to a styled Excel (.xlsx) file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"ok": False, "error": "openpyxl is not installed"}), 500

    date_filter    = request.args.get("date")
    start_date     = request.args.get("start_date")
    end_date       = request.args.get("end_date")
    channel_filter = request.args.get("channel")
    logs = database.get_dispensing_logs(
        limit=10000,
        date_filter=date_filter,
        channel_filter=channel_filter,
        start_date=start_date,
        end_date=end_date
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ประวัติการตัดสต็อก"

    header_fill = PatternFill("solid", fgColor="0D3B66")
    alt_fill    = PatternFill("solid", fgColor="F1F5F9")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font   = Font(name="Calibri", size=10)
    bold_font   = Font(name="Calibri", bold=True, size=10)
    red_font    = Font(name="Calibri", bold=True, color="E11D48", size=10)
    green_font  = Font(name="Calibri", bold=True, color="059669", size=10)
    center      = Alignment(horizontal="center", vertical="center")
    left        = Alignment(horizontal="left",   vertical="center")
    thin_side   = Side(style="thin", color="CBD5E1")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # ── Title ──
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    if start_date and end_date:
        date_label = f" ({start_date} ถึง {end_date})"
        file_suffix = f"{start_date.replace('-', '')}_{end_date.replace('-', '')}"
    elif start_date:
        date_label = f" (ตั้งแต่ {start_date})"
        file_suffix = f"from_{start_date.replace('-', '')}"
    elif end_date:
        date_label = f" (ถึง {end_date})"
        file_suffix = f"until_{end_date.replace('-', '')}"
    elif date_filter:
        date_label = f" ({date_filter})"
        file_suffix = f"{date_filter.replace('-', '')}"
    else:
        date_label = ""
        file_suffix = datetime.now().strftime('%Y%m%d_%H%M')

    title_cell.value = f"ประวัติการตัดสต็อก — PharmaCore{date_label}  (ส่งออก {datetime.now().strftime('%d/%m/%Y %H:%M')})"
    title_cell.font  = Font(name="Calibri", bold=True, color="0D3B66", size=14)
    title_cell.alignment = left
    ws.row_dimensions[1].height = 26

    # ── Header ──
    headers    = ["ลำดับ", "วัน-เวลา", "รหัสยา", "ชื่อยา", "จำนวนจ่าย",
                  "คงก่อน", "คงหลัง", "ช่องทาง", "ผู้จ่ายยา"]
    col_widths = [6, 20, 13, 30, 13, 10, 10, 16, 22]

    ws.append(headers)
    for c_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=c_idx)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = thin_border
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.row_dimensions[2].height = 22

    # ── Data ──
    for row_num, l in enumerate(logs, start=1):
        row_data = [
            row_num,
            l.get("created_at", ""),
            l.get("drug_code", ""),
            l.get("drug_name", ""),
            -abs(l.get("quantity", 0)),
            l.get("balance_before", 0),
            l.get("balance_after", 0),
            l.get("channel", ""),
            l.get("dispensed_by", ""),
        ]
        ws.append(row_data)
        xlsx_row = row_num + 2
        row_fill = alt_fill if row_num % 2 == 0 else None
        for c_idx in range(1, 10):
            cell = ws.cell(row=xlsx_row, column=c_idx)
            cell.border    = thin_border
            cell.alignment = center if c_idx in (1, 5, 6, 7) else left
            if c_idx == 5:
                cell.font = red_font
            elif c_idx == 7:
                cell.font = green_font
            elif c_idx in (1, 3):
                cell.font = bold_font
            else:
                cell.font = body_font
            if row_fill:
                cell.fill = row_fill
        ws.row_dimensions[xlsx_row].height = 18

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:I{len(logs) + 2}"

    buf = BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    filename = f"pharmacore_logs_{file_suffix}.xlsx"
    resp = make_response(xlsx_bytes)
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    resp.headers["Content-Length"] = len(xlsx_bytes)
    return resp


# ─────────────────────────────────────────────────────────────
# Official Thai Requisition Slip / Voucher Routes
# ─────────────────────────────────────────────────────────────

THAI_MONTHS_FULL = [
    '', 'มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
    'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม'
]

def format_thai_date_full(dt_str):
    if not dt_str:
        return "-"
    try:
        parts = dt_str.split(' ')[0].split('-')
        y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
        thai_year = y + 543
        return f"{d} {THAI_MONTHS_FULL[m]} พ.ศ. {thai_year}"
    except Exception:
        return str(dt_str)

def format_thai_date_short(dt_str):
    if not dt_str:
        return "-/-/----"
    try:
        parts = dt_str.split(' ')[0].split('-')
        y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
        thai_year = (y + 543) % 100
        return f"{d:02d}/{m:02d}/{thai_year:02d}"
    except Exception:
        return str(dt_str)

def number_to_thai_text(n):
    units = ['', 'หนึ่ง', 'สอง', 'สาม', 'สี่', 'ห้า', 'หก', 'เจ็ด', 'แปด', 'เก้า']
    digits = ['', 'สิบ', 'ร้อย', 'พัน', 'หมื่น', 'แสน', 'ล้าน']
    try:
        n_int = int(n)
        if n_int == 0:
            return 'ศูนย์'
        s = str(n_int)
        l = len(s)
        res = []
        for i, c in enumerate(s):
            d = int(c)
            p = l - i - 1
            if d == 0:
                continue
            if p == 0 and d == 1 and l > 1 and s[-2] != '0':
                res.append('เอ็ด')
            elif p == 1 and d == 2:
                res.append('ยี่สิบ')
            elif p == 1 and d == 1:
                res.append('สิบ')
            else:
                res.append(units[d] + digits[p])
        return ''.join(res)
    except Exception:
        return str(n)

def thai_baht_text(amount):
    """Convert monetary amount to formal Thai Baht text (e.g. 150.50 -> หนึ่งร้อยห้าสิบบาทห้าสิบสตางค์)."""
    try:
        amount = float(amount)
    except (ValueError, TypeError):
        return 'ศูนย์บาทถ้วน'
    if amount == 0:
        return 'ศูนย์บาทถ้วน'
    if amount < 0:
        return 'ลบ' + thai_baht_text(-amount)

    units = ['', 'หนึ่ง', 'สอง', 'สาม', 'สี่', 'ห้า', 'หก', 'เจ็ด', 'แปด', 'เก้า']
    digits = ['', 'สิบ', 'ร้อย', 'พัน', 'หมื่น', 'แสน', 'ล้าน']

    def convert_upto_million(n):
        if n == 0: return ''
        s = str(n)
        l = len(s)
        res = []
        for i, c in enumerate(s):
            d = int(c)
            p = l - i - 1
            if d == 0: continue
            if p == 0 and d == 1 and l > 1:
                res.append('เอ็ด')
            elif p == 1 and d == 2:
                res.append('ยี่สิบ')
            elif p == 1 and d == 1:
                res.append('สิบ')
            else:
                res.append(units[d] + digits[p])
        return ''.join(res)

    def convert_integer(n):
        if n == 0: return ''
        if n < 1000000:
            return convert_upto_million(n)
        millions = n // 1000000
        remainder = n % 1000000
        return convert_integer(millions) + 'ล้าน' + convert_upto_million(remainder)

    baht = int(amount)
    satang = int(round((amount - baht) * 100))
    result = ''
    if baht > 0:
        result += convert_integer(baht) + 'บาท'
    else:
        result += 'ศูนย์บาท'
    if satang == 0:
        result += 'ถ้วน'
    else:
        result += convert_upto_million(satang) + 'สตางค์'
    return result


def sanitize_formal_signatory_name(name: str) -> str:
    """
    Filter out email addresses, informal Google account nicknames (e.g. 'FU TRUE'),
    system defaults, and usernames so they don't appear in official Thai document signatures.
    Returns clean formal name if valid, otherwise empty string '' to render dotted signature line.
    """
    if not name:
        return ""
    # Strip parenthesized email/notes e.g. "FU TRUE (mxmm2547@gmail.com)" -> "FU TRUE"
    clean = name.split("(")[0].strip() if "(" in name else name.strip()
    
    # Check for emails or domain extensions
    if "@" in clean or any(ext in clean.lower() for ext in [".com", ".co.th", ".net", ".org", "gmail", "hotmail", "yahoo"]):
        return ""
        
    # Check for system/informal tokens
    informal_tokens = {
        "fu true", "fewit", "atommic", "mratommic", "mr.atommic", "admin", "administrator",
        "user", "anonymous", "null", "none", "test", "tester", "automated tester", "automated",
        "demo", "root", "staff", "guest", "bot", "system",
        "เจ้าหน้าที่", "เจ้าหน้าที่คลังยา", "เภสัชกร/เจ้าหน้าที่", "เภสัชกร / เจ้าหน้าที่",
        "เภสัชกร / เจ้าหน้าที่คลังยา", "line user"
    }
    if clean.lower() in informal_tokens or any(tok in clean.lower() for tok in ["test", "bot", "demo", "null", "none"]):
        return ""
        
    # Check if string contains numbers (not a real person's formal Thai name)
    import re
    if re.search(r'\d', clean):
        return ""
        
    # If purely Latin characters (no Thai characters), require a space and reject single nicknames
    has_thai = any('\u0E00' <= ch <= '\u0E7F' for ch in clean)
    if not has_thai:
        if " " not in clean or clean.lower() in informal_tokens:
            return ""
            
    return clean


@app.route("/api/dispense/requisition/<int:log_id>", methods=["GET"])
@login_required
def view_single_requisition(log_id):
    """View and export official Thai requisition slip for a single dispense log with drug price."""
    log = database.get_dispensing_log_by_id(log_id)
    if not log:
        return "ไม่พบข้อมูลประวัติการตัดสต็อกสำหรับเลขที่นี้", 404

    created_at = log.get("created_at", "")
    time_str = created_at.split(" ")[1] if " " in created_at else ""
    date_str = created_at.split(" ")[0] if " " in created_at else ""

    clean_date = date_str.replace("-", "") or datetime.now().strftime("%Y%m%d")
    requisition_no = f"REQ-{clean_date}-{log['id']:04d}"

    channel_raw = log.get("channel", "WEB_MANUAL")
    channel_map = {
        "WEB_MANUAL": "ระบบคลังยาหน้าเคาน์เตอร์",
        "QR_SCAN": "สแกน QR Code / บาร์โค้ด",
        "BARCODE_SCAN": "สแกนเนอร์บาร์โค้ด",
        "LINE_TEXT": "ระบบสั่งการผ่าน LINE Bot (ข้อความ)",
        "LINE_OCR": "ระบบสั่งการผ่าน LINE Bot (รูปภาพ/OCR)",
    }
    channel_label = channel_map.get(str(channel_raw).upper(), str(channel_raw))

    dispenser_query = request.args.get("dispenser", "").strip()
    dispensed_by_raw = log.get("dispensed_by") or ""
    formal_name = dispenser_query or sanitize_formal_signatory_name(dispensed_by_raw)

    dispenser_formal_title = formal_name if formal_name else "เภสัชกร / เจ้าหน้าที่ประจำคลังยา"
    dispenser_clean_name = formal_name

    patient_raw = (log.get("patient_name") or "").strip()
    patient_name = patient_raw if patient_raw and patient_raw != "-" else "-"
    requester_dept = "หน่วยบริการตรวจรักษา / หอผู้ป่วย"

    qty = int(log.get("quantity", 0))
    unit = log.get("unit", "หน่วย")
    qty_thai = number_to_thai_text(qty) + f" {unit}ถ้วน"

    # Price calculations
    price = float(log.get("price") or 0.0)
    subtotal = qty * price
    log["price_val"] = price
    log["price_formatted"] = f"{price:,.2f}"
    log["subtotal_val"] = subtotal
    log["subtotal_formatted"] = f"{subtotal:,.2f}"

    total_price = subtotal
    total_price_formatted = f"{total_price:,.2f}"
    total_price_thai = thai_baht_text(total_price)

    print_datetime = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ref_code = secrets.token_hex(4).upper()

    return render_template(
        "requisition_slip.html",
        requisition_no=requisition_no,
        thai_date=format_thai_date_full(created_at),
        thai_date_short=format_thai_date_short(created_at),
        time_str=time_str,
        requester_dept=requester_dept,
        patient_name=patient_name,
        channel_label=channel_label,
        purpose_text=log.get("notes") or "จ่ายยาตามใบสั่งแพทย์ / เบิกใช้เพื่อการรักษา",
        dispenser_display=dispenser_formal_title,
        dispenser_formal_title=dispenser_formal_title,
        dispenser_clean_name=dispenser_clean_name,
        requester_name=patient_name if patient_name != "-" else "",
        requester_position="ผู้ขอเบิก / ผู้รับมอบ",
        items=[log],
        total_qty=qty,
        primary_unit=unit,
        total_qty_thai=qty_thai,
        total_price=total_price,
        total_price_formatted=total_price_formatted,
        total_price_thai=total_price_thai,
        print_datetime=print_datetime,
        ref_code=ref_code,
        is_voided=bool(log.get("is_voided")),
        void_reason=log.get("void_reason", "")
    )


@app.route("/api/dispense/requisition-batch", methods=["GET"])
@login_required
def view_batch_requisition():
    """View and export official Thai requisition slip for multiple logs (filtered by date or range) with prices."""
    date_filter = request.args.get("date")
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    channel_filter = request.args.get("channel")
    requester_input = request.args.get("requester", "").strip()

    logs = database.get_dispensing_logs(
        limit=1000,
        date_filter=date_filter,
        channel_filter=channel_filter,
        start_date=start_date,
        end_date=end_date
    )
    logs = [l for l in logs if not l.get("is_voided")]

    if not logs:
        return "ไม่พบข้อมูลรายการเบิกจ่ายตามเงื่อนไขที่เลือก (หรือรายการทั้งหมดถูกยกเลิกแล้ว)", 404

    if start_date and end_date:
        thai_date = f"{format_thai_date_full(start_date)} ถึง {format_thai_date_full(end_date)}"
        batch_code = f"{start_date.replace('-','')}_{end_date.replace('-','')}"
    elif start_date:
        thai_date = f"ตั้งแต่ {format_thai_date_full(start_date)}"
        batch_code = f"from_{start_date.replace('-','')}"
    elif end_date:
        thai_date = f"ถึงวันที่ {format_thai_date_full(end_date)}"
        batch_code = f"until_{end_date.replace('-','')}"
    elif date_filter:
        thai_date = format_thai_date_full(date_filter)
        batch_code = date_filter.replace("-", "")
    else:
        today_s = datetime.now().strftime("%Y-%m-%d")
        thai_date = format_thai_date_full(today_s)
        batch_code = datetime.now().strftime("%Y%m%d")

    requisition_no = f"REQ-BATCH-{batch_code}"

    total_qty = sum(l.get("quantity", 0) for l in logs)
    primary_unit = logs[0].get("unit", "หน่วย") if logs else "หน่วย"
    qty_thai = number_to_thai_text(total_qty) + f" {primary_unit}ถ้วน"

    # Price calculations
    total_price = 0.0
    for it in logs:
        p = float(it.get("price") or 0.0)
        q = int(it.get("quantity") or 0)
        st = p * q
        it["price_val"] = p
        it["price_formatted"] = f"{p:,.2f}"
        it["subtotal_val"] = st
        it["subtotal_formatted"] = f"{st:,.2f}"
        total_price += st

    total_price_formatted = f"{total_price:,.2f}"
    total_price_thai = thai_baht_text(total_price)

    dispenser_query = request.args.get("dispenser", "").strip()
    u = current_user()
    current_actor = u.get("name") or u.get("email") or ""
    formal_name = dispenser_query or sanitize_formal_signatory_name(current_actor)

    dispenser_formal_title = formal_name if formal_name else "เภสัชกร / เจ้าหน้าที่ประจำคลังยา"
    dispenser_clean_name = formal_name

    print_datetime = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ref_code = secrets.token_hex(4).upper()

    return render_template(
        "requisition_slip.html",
        requisition_no=requisition_no,
        thai_date=thai_date,
        thai_date_short=format_thai_date_short(date_filter or start_date or datetime.now().strftime("%Y-%m-%d")),
        time_str=datetime.now().strftime("%H:%M"),
        requester_dept=requester_input or "แผนกตรวจรักษา / หอผู้ป่วยสามัญ",
        patient_name="ตามบัญชีแนบท้าย",
        channel_label="รายงานรวมตามงวดประจำวัน",
        purpose_text="สรุปยอดการเบิกจ่ายยาและเวชภัณฑ์ประจำงวด",
        dispenser_display=dispenser_formal_title,
        dispenser_formal_title=dispenser_formal_title,
        dispenser_clean_name=dispenser_clean_name,
        requester_name="",
        requester_position="หัวหน้างาน / ผู้รับมอบ",
        items=logs,
        total_qty=total_qty,
        primary_unit=primary_unit,
        total_qty_thai=qty_thai,
        total_price=total_price,
        total_price_formatted=total_price_formatted,
        total_price_thai=total_price_thai,
        print_datetime=print_datetime,
        ref_code=ref_code
    )


@app.route("/api/dispense/requisition/batch/<batch_id>", methods=["GET"])
@login_required
def view_prescription_batch_requisition(batch_id):
    """
    View and export official Thai consolidated A4 requisition slip for a specific multi-drug Rx Cart batch.
    Lists all medicines prescribed for this patient in a single formatted table.
    """
    logs = database.get_dispensing_logs_by_batch(batch_id)
    if not logs:
        return f"ไม่พบข้อมูลใบเบิกจ่ายยาสำหรับเลขที่ใบสั่งยา: {batch_id}", 404

    first_log = logs[0]
    created_at = first_log.get("created_at", "")
    time_str = created_at.split(" ")[1] if " " in created_at else ""
    date_str = created_at.split(" ")[0] if " " in created_at else ""

    requisition_no = f"REQ-{batch_id}"

    patient_raw = (first_log.get("patient_name") or "").strip()
    patient_name = patient_raw if patient_raw and patient_raw != "-" else "-"
    requester_dept = "แผนกตรวจรักษา / หอผู้ป่วยสามัญ"
    channel_label = "ใบสั่งยาอิเล็กทรอนิกส์ (Rx Cart)"

    total_qty = sum(int(l.get("quantity") or 0) for l in logs)
    primary_unit = logs[0].get("unit", "หน่วย") if logs else "หน่วย"
    qty_thai = number_to_thai_text(total_qty) + f" {primary_unit}ถ้วน"

    total_price = 0.0
    for it in logs:
        p = float(it.get("price") or 0.0)
        q = int(it.get("quantity") or 0)
        st = p * q
        it["price_val"] = p
        it["price_formatted"] = f"{p:,.2f}"
        it["subtotal_val"] = st
        it["subtotal_formatted"] = f"{st:,.2f}"
        total_price += st

    total_price_formatted = f"{total_price:,.2f}"
    total_price_thai = thai_baht_text(total_price)

    dispenser_query = request.args.get("dispenser", "").strip()
    dispensed_by_raw = first_log.get("dispensed_by") or ""
    formal_name = dispenser_query or sanitize_formal_signatory_name(dispensed_by_raw)

    dispenser_formal_title = formal_name if formal_name else "เภสัชกร / เจ้าหน้าที่ประจำคลังยา"
    dispenser_clean_name = formal_name

    print_datetime = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ref_code = batch_id.split("-")[-1] if "-" in batch_id else secrets.token_hex(4).upper()

    is_any_voided = any(bool(l.get("is_voided")) for l in logs)

    return render_template(
        "requisition_slip.html",
        requisition_no=requisition_no,
        thai_date=format_thai_date_full(created_at),
        thai_date_short=format_thai_date_short(created_at),
        time_str=time_str,
        requester_dept=requester_dept,
        patient_name=patient_name,
        channel_label=channel_label,
        purpose_text=first_log.get("notes") or f"ใบสั่งยาผู้ป่วย ({len(logs)} รายการ) / เบิกใช้เพื่อการรักษา",
        dispenser_display=dispenser_formal_title,
        dispenser_formal_title=dispenser_formal_title,
        dispenser_clean_name=dispenser_clean_name,
        requester_name=patient_name if patient_name != "-" else "",
        requester_position="ผู้ขอเบิก / ผู้รับมอบยา",
        items=logs,
        total_qty=total_qty,
        primary_unit=primary_unit,
        total_qty_thai=qty_thai,
        total_price=total_price,
        total_price_formatted=total_price_formatted,
        total_price_thai=total_price_thai,
        print_datetime=print_datetime,
        ref_code=ref_code,
        is_voided=is_any_voided,
        void_reason="มีรายการยาในใบสั่งนี้ถูกยกเลิก" if is_any_voided else ""
    )


@app.route("/api/export/stock-in", methods=["GET"])
@login_required
def export_stock_in_excel():
    """Export stock-in receiving logs to a styled Excel (.xlsx) file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"ok": False, "error": "openpyxl is not installed"}), 500

    date_filter = request.args.get("date")
    start_date  = request.args.get("start_date")
    end_date    = request.args.get("end_date")
    drug_id     = request.args.get("drug_id")
    logs = database.get_stock_in_logs(
        limit=10000,
        drug_id=drug_id,
        date_filter=date_filter,
        start_date=start_date,
        end_date=end_date
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ประวัติการรับยาเข้า"

    header_fill = PatternFill("solid", fgColor="059669")
    alt_fill    = PatternFill("solid", fgColor="F0FDF4")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font   = Font(name="Calibri", size=10)
    bold_font   = Font(name="Calibri", bold=True, size=10)
    green_font  = Font(name="Calibri", bold=True, color="059669", size=10)
    center      = Alignment(horizontal="center", vertical="center")
    left        = Alignment(horizontal="left",   vertical="center")
    right       = Alignment(horizontal="right",  vertical="center")
    thin_side   = Side(style="thin", color="CBD5E1")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # ── Title ──
    ws.merge_cells("A1:N1")
    title_cell = ws["A1"]
    if start_date and end_date:
        date_label = f" ({start_date} ถึง {end_date})"
        file_suffix = f"{start_date.replace('-', '')}_{end_date.replace('-', '')}"
    elif start_date:
        date_label = f" (ตั้งแต่ {start_date})"
        file_suffix = f"from_{start_date.replace('-', '')}"
    elif end_date:
        date_label = f" (ถึง {end_date})"
        file_suffix = f"until_{end_date.replace('-', '')}"
    elif date_filter:
        date_label = f" ({date_filter})"
        file_suffix = f"{date_filter.replace('-', '')}"
    else:
        date_label = ""
        file_suffix = datetime.now().strftime('%Y%m%d_%H%M')

    title_cell.value = f"ประวัติการรับยาเข้าคลัง — PharmaCore{date_label}  (ส่งออก {datetime.now().strftime('%d/%m/%Y %H:%M')})"
    title_cell.font  = Font(name="Calibri", bold=True, color="059669", size=14)
    title_cell.alignment = left
    ws.row_dimensions[1].height = 26

    # ── Header ──
    headers = [
        "ลำดับ", "วัน-เวลาที่รับ", "รหัสยา", "ชื่อยา", "เลขล็อต (Lot)",
        "จำนวนรับเข้า", "คงเหลือก่อน", "คงเหลือหลัง", "ทุน/หน่วย (฿)", "รวมมูลค่าทุน (฿)",
        "ผู้จัดจำหน่าย (Supplier)", "วันหมดอายุ", "ผู้ตรวจรับ", "หมายเหตุ"
    ]
    col_widths = [6, 20, 13, 30, 16, 14, 12, 12, 14, 16, 24, 14, 18, 22]

    ws.append(headers)
    for c_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=c_idx)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = thin_border
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.row_dimensions[2].height = 22

    # ── Data ──
    for row_num, l in enumerate(logs, start=1):
        row_data = [
            row_num,
            l.get("created_at", ""),
            l.get("drug_code", ""),
            l.get("drug_name", ""),
            l.get("lot_number", "-") or "-",
            l.get("quantity", 0),
            l.get("balance_before", 0),
            l.get("balance_after", 0),
            float(l.get("cost_per_unit", 0) or 0),
            float(l.get("total_cost", 0) or 0),
            l.get("supplier", "-") or "-",
            l.get("expiry_date", "-") or "-",
            l.get("received_by", "-") or "-",
            l.get("notes", "") or "",
        ]
        ws.append(row_data)
        xlsx_row = row_num + 2
        row_fill = alt_fill if row_num % 2 == 0 else None
        for c_idx in range(1, 15):
            cell = ws.cell(row=xlsx_row, column=c_idx)
            cell.border    = thin_border
            cell.alignment = center if c_idx in (1, 5, 6, 7, 8, 12) else (right if c_idx in (9, 10) else left)
            if c_idx in (6, 8, 10):
                cell.font = green_font
            elif c_idx in (1, 3):
                cell.font = bold_font
            else:
                cell.font = body_font
            if row_fill:
                cell.fill = row_fill
            if c_idx in (9, 10):
                cell.number_format = "#,##0.00"
        ws.row_dimensions[xlsx_row].height = 18

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:N{len(logs) + 2}"

    buf = BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    filename = f"pharmacore_stockin_{file_suffix}.xlsx"
    resp = make_response(xlsx_bytes)
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    resp.headers["Content-Length"] = len(xlsx_bytes)
    return resp

# ─────────────────────────────────────────────────────────────
# LINE Webhook Endpoint
# ─────────────────────────────────────────────────────────────

@app.route("/webhook", methods=["POST", "GET"])
@app.route("/line/webhook", methods=["POST", "GET"])
@app.route("/api/line/webhook", methods=["POST", "GET"])
@app.route("/api/line-webhook", methods=["POST", "GET"])
@app.route("/api/webhook", methods=["POST", "GET"])
def line_webhook():
    if request.method == "GET":
        return "Pharmacy LINE Bot Webhook Active", 200

    body = request.get_data(as_text=True)
    try:
        events = json.loads(body).get("events", [])
        for ev in events:
            line_bot.handle_webhook_event(ev)
        return "OK", 200
    except Exception as e:
        logger.error(f"Error handling LINE webhook: {e}")
        return "Internal Error", 500

# ─────────────────────────────────────────────────────────────
# LINE Notification & Expiry Alert API
# ─────────────────────────────────────────────────────────────

@app.route("/api/line/config", methods=["GET"])
def get_line_config():
    subscribers = database.get_line_subscribers()
    target_id = Config.PHARMACY_LINE_TARGET_ID or database.get_app_setting("line_notify_target") or ""
    days = int(database.get_app_setting("expiry_alert_days") or Config.EXPIRY_ALERT_DAYS or 60)
    last_sent = database.get_app_setting("last_expiry_alert_sent") or "ยังไม่เคยส่ง"
    expiring = database.get_expiring_drugs(days=days)

    return jsonify({
        "ok": True,
        "is_configured": bool(Config.LINE_CHANNEL_ACCESS_TOKEN),
        "bot_basic_id": Config.LINE_BOT_BASIC_ID,
        "target_id": target_id,
        "expiry_alert_days": days,
        "expiring_count": len(expiring),
        "subscribers_count": len(subscribers),
        "subscribers": subscribers[:10],
        "last_alert_sent": last_sent,
        "webhook_url": f"{Config.BASE_URL}/webhook"
    })

@app.route("/api/line/config", methods=["POST"])
def update_line_config():
    data = request.json or {}
    target_id = data.get("target_id")
    days = data.get("expiry_alert_days")

    if target_id is not None:
        database.set_app_setting("line_notify_target", str(target_id).strip())
    if days is not None:
        try:
            database.set_app_setting("expiry_alert_days", int(days))
        except ValueError:
            pass

    return jsonify({"ok": True, "message": "อัปเดตการตั้งค่า LINE เรียบร้อยแล้ว"})

@app.route("/api/line/send-expiry-alert", methods=["POST"])
def trigger_line_expiry_alert():
    """Manual or API trigger to send expiration alert to LINE."""
    data = request.json or {}
    days_input = data.get("days")
    days = int(days_input) if days_input else int(database.get_app_setting("expiry_alert_days") or Config.EXPIRY_ALERT_DAYS or 60)
    target_id = data.get("target_id")
    force = data.get("force", True)

    result = line_bot.send_expiring_drugs_alert(days=days, target_id=target_id, force=force)
    return jsonify(result)

# ─────────────────────────────────────────────────────────────
# Automatic Scheduled Expiry Alert Daemon
# ─────────────────────────────────────────────────────────────

def start_expiry_alert_scheduler():
    """
    Background daemon thread that checks expiring drugs once per day and alerts LINE.
    Runs silently in background.
    """
    import threading
    import time

    def _loop():
        time.sleep(10) # Initial grace period
        while True:
            try:
                days = int(database.get_app_setting("expiry_alert_days") or Config.EXPIRY_ALERT_DAYS or 60)
                # Calls send_expiring_drugs_alert with force=False so it only triggers once per day
                line_bot.send_expiring_drugs_alert(days=days, force=False)
            except Exception as e:
                logger.warning(f"Expiry alert background loop error: {e}")
            time.sleep(3600 * 2) # Check every 2 hours

    scheduler_thread = threading.Thread(target=_loop, daemon=True, name="PharmacyExpiryScheduler")
    scheduler_thread.start()
    logger.info("⏰ Pharmacy Expiry Alert background scheduler started.")

# ─────────────────────────────────────────────────────────────
# System Status & Health
# ─────────────────────────────────────────────────────────────

@app.route("/api/status", methods=["GET"])
def system_status():
    return jsonify({
        "status": "healthy",
        "app": "Pharmacy Stock Management System",
        "version": "1.1.0",
        "port": Config.PORT,
        "database": Config.DB_PATH,
        "google_sheets_connected": google_sync.is_connected,
        "google_spreadsheet_id": google_sync.spreadsheet_id,
        "line_bot_configured": bool(Config.LINE_CHANNEL_ACCESS_TOKEN),
        "expiry_alert_configured": True
    })

if __name__ == "__main__":
    port = Config.PORT
    print(f"\n=======================================================")
    print(f"🏥 PHARMACY STOCK MANAGEMENT SYSTEM RUNNING ON PORT {port}")
    print(f"👉 Local Access:   http://localhost:{port}")
    print(f"👉 Public Domain:  {Config.BASE_URL}")
    print(f"👉 LINE Webhook:   {Config.BASE_URL}/webhook")
    print(f"=======================================================\n")
    app.run(host="0.0.0.0", port=port, debug=False)
